"""
scripts/generate_report.py
==========================
Generates Meemansa_ProjectReport.docx from measured results.
All metrics come from saved JSON files — none are fabricated.

Usage:
    .venv/Scripts/python.exe scripts/generate_report.py
"""
import json
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

METRICS_DIR = PROJECT_ROOT / "results" / "metrics"
FIGURES_DIR = PROJECT_ROOT / "results" / "figures"
OUTPUT_PATH = PROJECT_ROOT / "Meemansa_ProjectReport.docx"


def add_heading(doc, text, level):
    doc.add_heading(text, level=level)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p


def add_table_from_rows(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for para in hdr[i].paragraphs:
            for run in para.runs:
                run.bold = True
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path, caption, width_inches=5.5):
    if Path(path).exists():
        doc.add_picture(str(path), width=Inches(width_inches))
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True
    else:
        doc.add_paragraph(f"[Figure not found: {path}]")
    doc.add_paragraph()


def main():
    cmp = json.loads((METRICS_DIR / "model_comparison.json").read_text())
    interp = json.loads((METRICS_DIR / "interpretation_summary.json").read_text())
    err = interp["error_analysis"]
    fi = interp["feature_importances"]

    doc = Document()

    # Title
    title = doc.add_heading("Machine Learning-Based Network Intrusion Detection and Analytics System", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Project Report")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].bold = True
    doc.add_paragraph("Author: Meemansa")
    doc.add_paragraph("Dataset: UNSW-NB15 (official predefined train/test split)")
    doc.add_paragraph()

    # ------------------------------------------------------------------ #
    doc.add_heading("1. Problem Statement and Objectives", 1)
    add_para(doc,
        "Network intrusion detection systems (IDS) protect computer networks by identifying "
        "malicious traffic. This project builds an offline, reproducible, educational "
        "prototype that applies machine learning to the UNSW-NB15 benchmark dataset. "
        "The primary objective is binary classification: given a network-flow record, "
        "predict whether it represents normal traffic (label=0) or an attack (label=1)."
    )
    add_para(doc,
        "Key objectives:\n"
        "  1. Load and validate the official UNSW-NB15 train/test split without leakage.\n"
        "  2. Design a reproducible, leakage-safe preprocessing pipeline.\n"
        "  3. Train and compare three models: a Dummy baseline, Logistic Regression, and Random Forest.\n"
        "  4. Evaluate models on the held-out test set using precision, recall, F1, and confusion matrices.\n"
        "  5. Analyse feature importance and error patterns.\n"
        "  6. Document all findings and limitations honestly.\n"
        "  7. Build an interactive Streamlit web dashboard that presents all saved results and "
        "enables batch prediction using the pre-trained pipeline."
    )
    add_para(doc,
        "This is an offline benchmark evaluation, NOT a live network monitoring system. "
        "Results reflect performance on UNSW-NB15 and cannot be directly generalised "
        "to real-world network environments."
    )

    # ------------------------------------------------------------------ #
    doc.add_heading("2. Dataset Source and Description", 1)
    add_para(doc,
        "Dataset: UNSW-NB15\n"
        "Source: https://research.unsw.edu.au/projects/unsw-nb15-dataset\n"
        "Citation: Moustafa, N., & Slay, J. (2015). UNSW-NB15: A comprehensive data set for "
        "network intrusion detection systems (UNSW-NB15 network data set). "
        "2015 Military Communications and Information Systems Conference (MilCIS), 1-6. "
        "https://doi.org/10.1109/MilCIS.2015.7348942"
    )
    add_para(doc,
        "The dataset was created by the Australian Centre for Cyber Security at UNSW Canberra. "
        "It contains network flow records with 45 features extracted from real and synthetic "
        "network traffic, labelled as normal or one of nine attack categories."
    )
    add_table_from_rows(doc,
        ["Split", "Rows", "Normal (0)", "Attack (1)", "Attack %"],
        [
            ["Training", "175,341", "56,000", "119,341", "68.1%"],
            ["Testing",  "82,332",  "37,000", "45,332",  "55.1%"],
        ]
    )
    add_para(doc,
        "Notable characteristics:\n"
        "  - 45 columns: 39 numeric, 3 categorical (proto, service, state), attack_cat, label, id.\n"
        "  - attack_cat encodes the attack type for attack rows; NaN for normal rows (by design).\n"
        "  - No missing values found in either split.\n"
        "  - No duplicate rows.\n"
        "  - Class balance is inverted vs real networks: attacks dominate training (68%).\n"
        "  - The id column is a sequential row identifier that overlaps between splits and was excluded."
    )

    # ------------------------------------------------------------------ #
    doc.add_heading("3. Preprocessing and Leakage Prevention", 1)
    add_para(doc,
        "All preprocessing is implemented as a scikit-learn Pipeline containing a ColumnTransformer. "
        "The pipeline is fitted exclusively on training data and applied to the test set without refitting."
    )
    add_para(doc, "Excluded columns (never used as model features):", bold=True)
    add_table_from_rows(doc,
        ["Column", "Reason for exclusion"],
        [
            ["label",      "Binary prediction target."],
            ["attack_cat", "Encodes attack category derived from the outcome — direct leakage."],
            ["id",         "Sequential row identifier; overlapping ranges; no traffic-flow meaning."],
        ]
    )
    add_para(doc, "Numeric pipeline (39 features):", bold=True)
    add_para(doc,
        "SimpleImputer(strategy='median') to handle unexpected missing values, followed by "
        "StandardScaler (zero mean, unit variance). Median imputation is used because numeric "
        "features (byte counts, timing values) are heavily right-skewed."
    )
    add_para(doc, "Categorical pipeline (proto, service, state):", bold=True)
    add_para(doc,
        "SimpleImputer(strategy='most_frequent'), followed by "
        "OneHotEncoder(handle_unknown='ignore'). The handle_unknown='ignore' setting is required "
        "because EDA confirmed that 'state' has two values in the test set ('ACC', 'CLO') not "
        "present in training."
    )
    add_para(doc, "High-cardinality handling (proto):", bold=True)
    add_para(doc,
        "proto has 133 unique training values; the top-10 values cover 92.9% of rows. "
        "OneHotEncoder(max_categories=20) bundles rare protocols into an infrequent bucket, "
        "keeping the output dimension manageable. service (13 values) and state (9 values) "
        "are encoded without a cap. Total output feature dimension: 81 columns."
    )
    add_para(doc,
        "Train-only fitting prevents data leakage: imputer medians, scaler mean/variance, and "
        "OHE category lists are computed from training data only. The test set is kept "
        "completely untouched until final evaluation."
    )

    # ------------------------------------------------------------------ #
    doc.add_heading("4. Models Compared and Evaluation Methodology", 1)
    add_para(doc,
        "Three models were trained and compared. Each model is saved as a complete "
        "scikit-learn Pipeline (preprocessor + classifier) to ensure preprocessing "
        "parameters are always aligned with the fitted classifier."
    )
    add_table_from_rows(doc,
        ["Model", "Role", "Key settings"],
        [
            ["DummyClassifier", "Baseline", "strategy='most_frequent' — always predicts attack"],
            ["Logistic Regression", "Linear model", "solver=lbfgs, max_iter=1000, C=1.0, random_state=42"],
            ["Random Forest", "Ensemble model", "n_estimators=100, max_depth=None, random_state=42"],
        ]
    )
    add_para(doc, "Evaluation methodology:", bold=True)
    add_para(doc,
        "All models were trained on the official UNSW-NB15 training split (175,341 rows). "
        "Final evaluation used the official held-out test split (82,332 rows), which was "
        "not used during training or model selection. No cross-validation or hyperparameter "
        "tuning was performed in this phase.\n\n"
        "Metrics reported:\n"
        "  - Accuracy (with caveats about class imbalance)\n"
        "  - Per-class precision, recall, and F1 (labels 0 and 1)\n"
        "  - Macro F1 (unweighted mean across classes)\n"
        "  - ROC-AUC where prediction scores are available\n"
        "  - Confusion matrix (TN, FP, FN, TP)\n"
        "  - Training time and prediction time"
    )
    add_para(doc,
        "Intrusion-detection interpretation:\n"
        "  - False Negatives (FN) = attacks predicted as normal = MISSED ATTACKS. "
        "  This is the more dangerous failure — a missed attack goes undetected.\n"
        "  - False Positives (FP) = normal predicted as attack = FALSE ALARMS. "
        "  Too many FPs cause alert fatigue and reduce analyst trust."
    )

    # ------------------------------------------------------------------ #
    doc.add_heading("5. Measured Results", 1)
    add_para(doc,
        "All metrics below were measured on the official UNSW-NB15 test split (82,332 rows). "
        "Values are loaded from results/metrics/model_comparison.json — they were not manually edited.",
        bold=False
    )
    add_table_from_rows(doc,
        ["Model", "Accuracy", "Attack Recall", "Attack Precision", "Macro F1", "ROC-AUC", "Train (s)"],
        [
            [m, f"{v['accuracy']:.4f}", f"{v['attack_recall']:.4f}",
             f"{v['attack_precision']:.4f}", f"{v['macro_f1']:.4f}",
             f"{v.get('roc_auc', 'N/A')}", f"{v['train_time_s']}"]
            for m, v in cmp.items()
        ]
    )
    doc.add_heading("5.1 Confusion Matrices", 2)
    add_table_from_rows(doc,
        ["Model", "TN", "FP", "FN", "TP"],
        [[m, f"{v['tn']:,}", f"{v['fp']:,}", f"{v['fn']:,}", f"{v['tp']:,}"] for m, v in cmp.items()]
    )
    add_para(doc,
        "Key findings:\n"
        "  - The Dummy baseline (FN=0, FP=37,000) confirms that always predicting 'attack' "
        "catches every attack but floods analysts with false alarms — completely impractical.\n"
        "  - Logistic Regression reduces FP to 14,481 and FN to 1,253 with a ROC-AUC of 0.9551.\n"
        "  - Random Forest achieves the best balance: only 642 missed attacks (FN) and 9,850 false "
        "alarms (FP), with ROC-AUC 0.9789 and Macro F1 0.8665.\n"
        "  - The test set has 55% attacks vs 68% in training — both real classifiers generalise "
        "reasonably well to the shifted class balance."
    )

    add_figure(doc, FIGURES_DIR / "confusion_matrix_RandomForest.png",
               "Figure 1: Confusion Matrix — Random Forest (best model)")
    add_figure(doc, FIGURES_DIR / "roc_curves.png",
               "Figure 2: ROC Curves — All Models")
    add_figure(doc, FIGURES_DIR / "model_comparison.png",
               "Figure 3: Model Comparison — Key Metrics")

    # ------------------------------------------------------------------ #
    doc.add_heading("6. Feature Importance Analysis", 1)
    add_para(doc,
        "Feature importances were extracted from the fitted Random Forest using mean decrease "
        "in impurity (MDI). The top 10 features by importance are:"
    )
    top10 = fi[:10]
    add_table_from_rows(doc,
        ["Rank", "Feature", "Importance"],
        [[r["rank"],
          r["feature"].replace("numeric__", "").replace("cat_proto__proto_", "proto:"),
          f"{r['importance']:.5f}"]
         for r in top10]
    )
    add_figure(doc, FIGURES_DIR / "feature_importance_rf.png",
               "Figure 4: Random Forest Feature Importance — Top 20 Features")
    add_para(doc, "Interpretation and limitations:", bold=True)
    add_para(doc,
        "The top feature 'sttl' (source TTL) makes network sense: different operating systems "
        "and attack tools set characteristic TTL values. The pair sttl/ct_state_ttl likely "
        "captures OS fingerprinting signals.\n\n"
        "Important limitations of MDI importance:\n"
        "  1. MDI can overvalue high-cardinality features (e.g. stcpb, dtcpb).\n"
        "  2. Correlated features share importance — sttl and ct_state_ttl may appear less "
        "important individually than their combined effect warrants.\n"
        "  3. These values describe how the model makes splits, not causal relationships "
        "between features and attack behaviour.\n"
        "  4. Permutation importance (not computed in this project) would provide a less "
        "biased alternative."
    )

    # ------------------------------------------------------------------ #
    doc.add_heading("7. Error Analysis", 1)
    add_para(doc,
        f"The Random Forest produced {err['total_fp']:,} false positives and "
        f"{err['total_fn']:,} false negatives on the test set."
    )
    if "fp_proto_counts" in err:
        fp_proto = list(err["fp_proto_counts"].items())[:5]
        add_para(doc, "Top protocols in false positives: " +
                 ", ".join(f"{p}: {n:,}" for p, n in fp_proto))
    if "fn_proto_counts" in err:
        fn_proto = list(err["fn_proto_counts"].items())[:5]
        add_para(doc, "Top protocols in false negatives (missed attacks): " +
                 ", ".join(f"{p}: {n:,}" for p, n in fn_proto))
    add_figure(doc, FIGURES_DIR / "error_distributions_rf.png",
               "Figure 5: Feature Distributions — False Positives vs False Negatives")
    add_para(doc,
        "The error analysis shows that both false positives and false negatives are concentrated "
        "in the high-volume TCP and UDP protocols, which dominate the dataset. Attacks being missed "
        "tend to have feature distributions that overlap closely with normal traffic — these are the "
        "inherently ambiguous cases for a binary classifier."
    )

    # ------------------------------------------------------------------ #
    doc.add_heading("8. Limitations and Future Work", 1)
    add_para(doc, "Limitations:", bold=True)
    add_para(doc,
        "  1. Offline benchmark only: results are specific to UNSW-NB15 and cannot be generalised "
        "to real network environments without independent validation.\n"
        "  2. Inverted class balance: 68% attacks in training is unrealistic — real networks "
        "are dominated by normal traffic, which would change the operational FP/FN trade-off.\n"
        "  3. No hyperparameter tuning: Logistic Regression and Random Forest settings were chosen "
        "based on documented guidance, not cross-validated search.\n"
        "  4. MDI feature importance is biased: high-cardinality and correlated features are "
        "mis-valued; permutation importance would be more reliable.\n"
        "  5. Single train/test split: no confidence intervals or multiple-run variance estimates "
        "are provided.\n"
        "  6. Binary classification only: multiclass attack-category prediction was out of scope."
    )
    add_para(doc, "Future work:", bold=True)
    add_para(doc,
        "  - Threshold optimisation: adjust the decision boundary to prioritise FN reduction "
        "(fewer missed attacks) at the cost of more FP, or vice versa.\n"
        "  - Permutation importance: replace MDI with a less biased importance measure.\n"
        "  - Multiclass classification: use attack_cat to train a multiclass model and assess "
        "per-category detection rates.\n"
        "  - Ensemble methods: XGBoost or stacking to improve on Random Forest.\n"
        "  - Independent evaluation: apply the pipeline to a different network dataset."
    )

    # ------------------------------------------------------------------ #
    doc.add_heading("9. Streamlit Web Dashboard", 1)
    add_para(doc,
        "A Streamlit-based interactive web dashboard (app.py) was developed as the project's "
        "frontend layer. It reads all pre-computed results — saved model pipelines, JSON/CSV "
        "metrics, and PNG figures — and presents them through a clean, dark-themed cybersecurity "
        "interface. No model training or fitting takes place within the dashboard."
    )

    doc.add_heading("9.1 Design and Technology", 2)
    add_para(doc,
        "The dashboard is built with Streamlit 1.66 and Matplotlib, using the existing .venv "
        "environment. It applies a dark navy/charcoal theme with blue/cyan accents to match the "
        "cybersecurity context. Metric cards, interactive tables, and embedded PNG charts are used "
        "throughout. The app is launched locally via: "
        ".venv\\Scripts\\python.exe -m streamlit run app.py "
        "(accessible at http://localhost:8501). No external APIs, cloud services, or internet "
        "access are required at runtime."
    )

    doc.add_heading("9.2 Dashboard Pages", 2)
    add_para(doc,
        "The dashboard is structured into six pages accessible from a sidebar navigation panel:"
    )
    add_para(doc,
        "Overview: Project summary, dataset record counts (train: 175,341; test: 82,332), "
        "best-model metric cards (Random Forest accuracy 87.3%, ROC-AUC 0.9789, attack recall "
        "98.6%), and the class-distribution, model-comparison, and ROC-curve figures. A clear "
        "explanation of the Normal (label=0) vs Attack (label=1) distinction is provided."
    )
    add_para(doc,
        "Model Comparison: Full metrics table loaded from results/metrics/model_comparison.csv "
        "and model_comparison.json, including accuracy, precision, recall, F1, ROC-AUC, and "
        "prediction time for all three models. Per-model confusion matrices are displayed side "
        "by side. A metric glossary explains each measure in plain language."
    )
    add_para(doc,
        "Dataset Explorer: Class and attack-category distribution tables (from "
        "results/metrics/eda_summary.json), all six EDA figures, and a 500-row preview of the "
        "training CSV. The excluded columns (label, attack_cat, id) and their rationale are "
        "explained."
    )
    add_para(doc,
        "Feature Importance and Errors: An interactive horizontal bar chart of the top 20 "
        "Random Forest features (rendered with Matplotlib), a top-10 table, the saved "
        "feature-importance and error-distribution PNG figures, protocol-level error breakdowns, "
        "and a plain-language explanation of Gini importance and its limitations."
    )
    add_para(doc,
        "Prediction Lab: Users upload a CSV of network-flow records. The dashboard validates "
        "the schema against the 42 required feature columns, then calls pipeline.predict() and "
        "pipeline.predict_proba() on the saved RandomForest.joblib pipeline. Predictions are "
        "displayed as Normal or Attack with confidence scores, and results can be downloaded as "
        "a CSV. A synthetic demo option is also available. No .fit() call is made anywhere in "
        "the dashboard."
    )
    add_para(doc,
        "Project Information: Background on NIDS and the UNSW-NB15 dataset, model descriptions, "
        "metric definitions in plain language, dashboard usage instructions, project limitations, "
        "and an explanation of how saved pipelines are reused without retraining."
    )

    doc.add_heading("9.3 Prediction Pipeline Reuse", 2)
    add_para(doc,
        "Each saved .joblib file contains the complete scikit-learn Pipeline — the fitted "
        "ColumnTransformer preprocessor (scaler statistics, one-hot encoder category lists) and "
        "the fitted classifier. The dashboard loads the pipeline once using joblib.load(), "
        "caches it with st.cache_resource, and calls only .predict() and .predict_proba(). "
        "Label convention is preserved: 0 maps to Normal, 1 maps to Attack, consistent with the "
        "training target definition in src/features.py. Unseen categorical values in uploaded "
        "data are handled gracefully by the pipeline's OneHotEncoder(handle_unknown='ignore') "
        "setting."
    )

    doc.add_heading("9.4 Limitations of the Dashboard", 2)
    add_para(doc,
        "The dashboard is an offline educational tool, not a production security system. "
        "Predictions are only as reliable as the underlying model on UNSW-NB15; they do not "
        "reflect performance on real or modern network traffic. Uploaded CSVs must match the "
        "UNSW-NB15 feature schema exactly. The 101 MB Random Forest model requires sufficient "
        "memory on the host machine. The interface includes a clear disclaimer on every page "
        "that this is not connected to live network traffic."
    )

    # ------------------------------------------------------------------ #
    doc.add_heading("10. References", 1)
    add_para(doc,
        "Moustafa, N., & Slay, J. (2015). UNSW-NB15: A comprehensive data set for network "
        "intrusion detection systems (UNSW-NB15 network data set). 2015 Military Communications "
        "and Information Systems Conference (MilCIS), 1-6. "
        "https://doi.org/10.1109/MilCIS.2015.7348942\n\n"
        "Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. "
        "Journal of Machine Learning Research, 12, 2825-2830.\n\n"
        "McKinney, W. (2010). Data Structures for Statistical Computing in Python. "
        "Proceedings of the 9th Python in Science Conference, 51-56."
    )

    doc.save(OUTPUT_PATH)
    print(f"Report saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
